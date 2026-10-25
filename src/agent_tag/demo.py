"""Jeu de démonstration TEST : base séparée, données FICTIVES du classeur de référence.

Crée (ou recrée) la base indiquée, applique les migrations, initialise un environnement TEST
libellé « DÉMONSTRATION », crée deux utilisateurs de test, importe le référentiel et les données
journalières du 12/10/2026 (absences attestées, choix provisoire D18), puis installe le profil
PROVISOIRE A. Rien n'est validé ; aucune donnée réelle n'est utilisée.
"""

from __future__ import annotations

from pathlib import Path

import psycopg
from psycopg import conninfo, sql

from agent_tag.imports import import_journalier, import_referentiel
from agent_tag.persistence import migrations
from agent_tag.services import configuration

CLASSEUR = (Path(__file__).resolve().parents[2] / "donnees" / "reference"
            / "Base_Allocation_Production_150_Ouvriers_v2.xlsx")
UTILISATEURS = [
    ("demo.responsable", "Responsable de démonstration", "TEST_TOUS_DROITS"),
    ("demo.lecteur", "Lecteur de démonstration (sans droit de décision)", "LECTURE_TEST"),
]


def initialiser(owner_url: str, app_url: str, nom_base: str, dossier_archives: Path,
                profil: str = "A") -> dict:
    if "demo" not in nom_base:
        raise ValueError("par sécurité, le nom de la base de démonstration doit contenir « demo »")
    with psycopg.connect(owner_url, autocommit=True) as admin:
        admin.execute(sql.SQL("DROP DATABASE IF EXISTS {} WITH (FORCE)").format(sql.Identifier(nom_base)))
        admin.execute(sql.SQL("CREATE DATABASE {} TEMPLATE template0 ENCODING 'UTF8'")
                      .format(sql.Identifier(nom_base)))
        admin.execute(sql.SQL("REVOKE ALL ON DATABASE {} FROM PUBLIC").format(sql.Identifier(nom_base)))
        admin.execute(sql.SQL("GRANT CONNECT ON DATABASE {} TO agent_tag_app")
                      .format(sql.Identifier(nom_base)))
    owner_demo = conninfo.make_conninfo(owner_url, dbname=nom_base)
    app_demo = conninfo.make_conninfo(app_url, dbname=nom_base)
    migrations.migrer(owner_demo)
    with psycopg.connect(owner_demo) as conn:
        conn.execute("INSERT INTO environment (usage, libelle) VALUES ('TEST', %s)",
                     ("DÉMONSTRATION — données fictives",))
        with conn.cursor() as cur:
            cur.executemany("INSERT INTO app_user (identifiant, nom, role) VALUES (%s, %s, %s)",
                            UTILISATEURS)
    ref = import_referentiel.importer(app_demo, CLASSEUR, "demo.responsable", dossier_archives)
    jour = import_journalier.importer(app_demo, CLASSEUR, "demo.responsable", dossier_archives,
                                      attester_absences=True)
    version = configuration.installer_profil_provisoire(owner_demo, profil, "demo.responsable")
    from urllib.parse import urlsplit, urlunsplit
    morceaux = urlsplit(app_url) if "://" in app_url else None
    url_affichee = urlunsplit(morceaux._replace(path="/" + nom_base)) if morceaux else app_demo
    return {"base": nom_base, "url_application": url_affichee, "referentiel": ref.statut,
            "journalier": jour.statut, "configuration": version}


def simuler_reimport(app_url: str, auteur: str, id_ouvrier: str, dossier_archives: Path,
                     dossier_travail: Path) -> "object":
    """DÉMONSTRATION TEST : produit une copie du classeur de référence où l'ouvrier indiqué,
    présent le 12/10/2026, devient absent (affectation « Absent » + ligne d'absence d'un jour),
    puis l'importe comme nouveau lot journalier. Sert à montrer les effets d'un réimport
    (calculs à recalculer, remplacements à revoir). Aucune donnée réelle n'est produite."""
    from datetime import datetime

    from openpyxl import load_workbook

    classeur = load_workbook(CLASSEUR, data_only=True)   # valeurs figées (dates calculées)
    aff = classeur["Affectations"]
    entetes = [c.value for c in aff[1]]
    trouve = None
    for ligne in aff.iter_rows(min_row=2):
        valeurs = dict(zip(entetes, (c.value for c in ligne)))
        if valeurs["ID_Ouvrier"] == id_ouvrier:
            if valeurs["Statut_affectation"] != "Présent":
                raise ValueError(f"{id_ouvrier} n'est pas présent dans le classeur de référence")
            ligne[entetes.index("Statut_affectation")].value = "Absent"
            trouve = valeurs
    if trouve is None:
        raise ValueError(f"ouvrier {id_ouvrier} absent de la feuille Affectations")
    classeur["Absences"].append([datetime(2026, 10, 12), id_ouvrier, trouve["Nom"],
                                 trouve["Chaîne"], trouve["Poste"], "Absent",
                                 "Absence simulée (démonstration TEST)", 1])
    dossier_travail.mkdir(parents=True, exist_ok=True)
    chemin = dossier_travail / f"demo_reimport_{id_ouvrier}.xlsx"
    classeur.save(chemin)
    return import_journalier.importer(app_url, chemin, auteur, dossier_archives,
                                      attester_absences=True)
