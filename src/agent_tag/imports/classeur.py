"""Lecture du classeur Excel : feuilles, en-têtes et lignes, sans interprétation métier."""

from __future__ import annotations

import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from agent_tag.imports import rapport as r

# Colonnes attendues par feuille (noms exacts du classeur de référence).
COLONNES = {
    "Ouvriers": ["ID_Ouvrier", "Nom", "Chaîne_actuelle", "Produit", "Poste_principal",
                 "Statut", "Expérience_poste_ans", "Horaire"],
    "Competences": ["ID_Ouvrier", "Nom", "Compétence", "Type_compétence",
                    "Années_expérience", "Niveau", "Certification"],
    "Chaines": ["ID_Chaîne", "Produit", "Effectif_nécessaire", "Effectif_planifié",
                "Priorité", "Compétences_critiques", "Capacité_journalière_unités"],
    "Besoins_Postes": ["Chaîne", "Produit", "Poste", "Nombre_requis",
                       "Compétence_requise", "Niveau_minimum"],
    "Absences": ["Date", "ID_Ouvrier", "Nom", "Chaîne", "Poste_principal",
                 "Statut_absence", "Motif", "Durée_jours"],
    "Affectations": ["Date", "ID_Ouvrier", "Nom", "Chaîne", "Produit", "Poste",
                     "Horaire", "Statut_affectation"],
}
FEUILLES_REFERENTIEL = ("Chaines", "Ouvriers", "Competences", "Besoins_Postes")
FEUILLES_JOURNALIERES = ("Absences", "Affectations")
FEUILLES_INFORMATIVES = ("Scenario_Test", "README")


def nfc(texte: str) -> str:
    return unicodedata.normalize("NFC", texte)


@dataclass
class Ligne:
    numero: int          # numéro de ligne Excel (l'en-tête est la ligne 1)
    valeurs: dict


@dataclass
class Feuille:
    nom: str
    lignes: list[Ligne]


class ClasseurIllisible(Exception):
    pass


def lire(chemin: Path, rapport: r.Rapport,
         obligatoires=FEUILLES_REFERENTIEL,
         controlees=FEUILLES_JOURNALIERES,
         message_controlees="structure contrôlée ; données importées par importer-journalier"
         ) -> dict[str, Feuille]:
    """Lit les feuilles d'un classeur.

    - `obligatoires` : absence ou structure incorrecte = anomalie BLOQUANTE ; lignes retournées ;
    - `controlees` : structure contrôlée (AVERTISSEMENT si incorrecte), lignes lues mais non
      importées par cet import ;
    - toute autre feuille connue présente est signalée comme non importée.
    Retourne les feuilles dont toutes les colonnes attendues sont présentes.
    """
    try:
        classeur = load_workbook(chemin, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ClasseurIllisible(f"classeur illisible : {exc}") from exc

    try:
        presentes = {nfc(n.strip()): n for n in classeur.sheetnames}
        lues: dict[str, Feuille] = {}

        for nom in (*COLONNES, *FEUILLES_INFORMATIVES):
            if nom in presentes and nom not in obligatoires and nom not in controlees:
                rapport.ajouter(r.INFORMATION, "FEUILLE_NON_IMPORTEE",
                                "feuille non concernée par cet import", feuille=nom)

        for nom in (*obligatoires, *controlees):
            attendues = COLONNES[nom]
            controlee = nom not in obligatoires
            gravite = r.AVERTISSEMENT if controlee else r.BLOQUANTE
            if nom not in presentes:
                rapport.ajouter(gravite, "FEUILLE_MANQUANTE",
                                "feuille absente" if controlee else "feuille obligatoire absente",
                                feuille=nom)
                continue

            lignes_brutes = list(classeur[presentes[nom]].iter_rows(values_only=True))
            entetes = [nfc(str(v).strip()) if v is not None else "" for v in
                       (lignes_brutes[0] if lignes_brutes else ())]
            manquantes = [c for c in attendues if c not in entetes]
            for c in manquantes:
                rapport.ajouter(gravite, "COLONNE_MANQUANTE", "colonne attendue absente",
                                feuille=nom, colonne=c)
            for c in entetes:
                if c and c not in attendues:
                    rapport.ajouter(r.AVERTISSEMENT, "COLONNE_INATTENDUE",
                                    "colonne non attendue, ignorée", feuille=nom, colonne=c)
            doublons = {c for c in entetes if c and entetes.count(c) > 1}
            for c in doublons:
                rapport.ajouter(gravite, "COLONNE_DUPLIQUEE", "en-tête présent plusieurs fois",
                                feuille=nom, colonne=c)
            if manquantes or doublons:
                continue

            index = {c: entetes.index(c) for c in attendues}
            lignes = []
            for numero, brute in enumerate(lignes_brutes[1:], start=2):
                if brute is None or all(v is None or (isinstance(v, str) and not v.strip())
                                        for v in brute):
                    continue
                valeurs = {c: (brute[i] if i < len(brute) else None) for c, i in index.items()}
                lignes.append(Ligne(numero, valeurs))
            lues[nom] = Feuille(nom, lignes)
            rapport.stat(nom).lues = len(lignes)
            if controlee:
                rapport.ajouter(r.INFORMATION, "FEUILLE_NON_IMPORTEE", message_controlees,
                                feuille=nom)
        return lues
    finally:
        classeur.close()
