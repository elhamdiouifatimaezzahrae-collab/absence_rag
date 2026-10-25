"""Incrément 4 : calculs persistés, propositions, décisions, réservations, réimport, concurrence.

Données : classeur de référence (fictif) importé dans une base jetable ; profil PROVISOIRE A.
Après chaque opération, le vérificateur d'invariants doit être vide.
"""

import datetime as dt
import threading
from pathlib import Path

import psycopg
import pytest
from openpyxl import load_workbook

from agent_tag.imports import import_journalier
from agent_tag.services import allocation as s
from agent_tag.services import configuration, integrite

pytestmark = pytest.mark.db

CLASSEUR = (Path(__file__).resolve().parents[2] / "donnees" / "reference"
            / "Base_Allocation_Production_150_Ouvriers_v2.xlsx")
J = dt.date(2026, 10, 12)


@pytest.fixture
def base(importer, base_import, tmp_path):
    if not CLASSEUR.is_file():
        pytest.skip("classeur de référence absent")
    assert importer(CLASSEUR).statut == "APPLIQUE"
    assert import_journalier.importer(base_import.app_url, CLASSEUR, "testeur",
                                      tmp_path / "arch", True).statut == "APPLIQUE"
    with base_import.owner() as c:
        c.execute("INSERT INTO app_user (identifiant, nom, role) "
                  "VALUES ('lecteur', 'Lecteur', 'LECTURE_TEST')")
    configuration.installer_profil_provisoire(base_import.owner_url, "A", "testeur")
    base_import.url = base_import.app_url
    base_import.archives = tmp_path / "arch"
    return base_import


def q(base, sql, *params):
    with psycopg.connect(base.url) as c:
        return c.execute(sql, params).fetchall()


def sain(base):
    with psycopg.connect(base.url) as c:
        violations = integrite.verifier(c)
    assert violations == {}, violations


def actives(base, id_run=None):
    return {absent: (pid, o) for pid, absent, o in q(
        base, "SELECT id, id_ouvrier_absent, id_ouvrier FROM proposal WHERE statut = 'ACTIVE'"
        + (" AND id_run = %s" if id_run else ""), *([id_run] if id_run else []))}


# --- Lancement ------------------------------------------------------------------------------

def test_lancer_enregistre_le_calcul_sans_aucune_affectation(base):
    id_run = s.lancer(base.url, J, "L1", "testeur")
    assert q(base, "SELECT statut, version FROM allocation_run WHERE id = %s", id_run)[0][0] == "OUVERT"
    assert q(base, "SELECT numero, declencheur, couverture_sequentielle, couverture_maximale "
                   "FROM allocation_iteration") == [(1, "INITIAL", 8, 8)]
    statuts = dict(q(base, "SELECT statut_cycle, count(*) FROM replacement_request GROUP BY 1"))
    assert statuts == {"EN_ATTENTE": 8, "NON_POURVU": 1, "BLOQUE": 1}
    assert q(base, "SELECT count(*) FROM worker_day_slot WHERE type = 'RESERVE' "
                   "AND statut = 'COURANT'") == [(8,)]
    # Une proposition n'est jamais une affectation confirmée.
    assert q(base, "SELECT count(*) FROM replacement_assignment") == [(0,)]
    assert q(base, "SELECT count(*) FROM allocation_decision") == [(0,)]
    # Évaluations et justifications conservées.
    assert q(base, "SELECT count(*) FROM candidate_evaluation WHERE explication = ''") == [(0,)]
    assert q(base, "SELECT cause_statut FROM replacement_request WHERE id_ouvrier_absent = 'O004'"
             )[0][0].startswith("EXIGENCE_ABSENTE")
    sain(base)


def test_lancer_refuse_si_deja_ouvert_ou_donnees_incompletes(base):
    s.lancer(base.url, J, "L1", "testeur")
    with pytest.raises(s.RefusMetier) as e:
        s.lancer(base.url, J, "L1", "testeur")
    assert e.value.code == "RUN_DEJA_OUVERT"
    with pytest.raises(s.RefusMetier) as e:
        s.lancer(base.url, J + dt.timedelta(days=1), "L1", "testeur")
    assert e.value.code == "DONNEES_INCOMPLETES"


def test_lancer_refuse_sans_configuration_active(importer, base_import, tmp_path):
    importer(CLASSEUR)
    import_journalier.importer(base_import.app_url, CLASSEUR, "testeur", tmp_path / "a", True)
    with pytest.raises(s.RefusMetier) as e:
        s.lancer(base_import.app_url, J, "L1", "testeur")
    assert e.value.code == "CONFIGURATION_ABSENTE"


# --- Confirmation et refus ------------------------------------------------------------------

def test_accepter_cree_le_remplacement_et_engage_l_ouvrier(base):
    s.lancer(base.url, J, "L1", "testeur")
    pid, ouvrier = actives(base)["O002"]
    id_rempl = s.accepter(base.url, pid, "testeur")
    assert q(base, "SELECT id_ouvrier, statut FROM replacement_assignment WHERE id = %s",
             id_rempl) == [(ouvrier, "ACTIF")]
    assert q(base, "SELECT type, statut FROM worker_day_slot WHERE id_proposal = %s", pid) == \
        [("ENGAGE", "COURANT")]
    assert q(base, "SELECT statut_cycle FROM replacement_request WHERE id_ouvrier_absent = 'O002'"
             ) == [("POURVU",)]
    # Décision tracée et marquée TEST / configuration non validée.
    assert q(base, "SELECT type, usage_environnement, configuration_validee "
                   "FROM allocation_decision") == [("ACCEPTER", "TEST", False)]
    sain(base)


def test_double_acceptation_refusee(base):
    s.lancer(base.url, J, "L1", "testeur")
    pid, _ = actives(base)["O002"]
    s.accepter(base.url, pid, "testeur")
    with pytest.raises(s.RefusMetier) as e:
        s.accepter(base.url, pid, "testeur")
    assert e.value.code == "PROPOSITION_NON_ACTIVE"
    assert q(base, "SELECT count(*) FROM replacement_assignment") == [(1,)]


def test_refuser_exige_un_motif_puis_recalcule_sans_l_ouvrier_refuse(base):
    id_run = s.lancer(base.url, J, "L1", "testeur")
    pid, refuse = actives(base)["O010"]
    with pytest.raises(s.RefusMetier) as e:
        s.refuser(base.url, pid, "testeur", "  ")
    assert e.value.code == "MOTIF_REQUIS"
    s.refuser(base.url, pid, "testeur", "indisponible en réalité")
    nouvelle = actives(base)["O010"]
    assert nouvelle[1] != refuse
    assert q(base, "SELECT max(numero) FROM allocation_iteration WHERE id_run = %s", id_run) == [(2,)]
    exclu = q(base, "SELECT premier_echec FROM candidate_evaluation e JOIN replacement_request r "
                    "ON r.id = e.id_besoin WHERE r.id_ouvrier_absent = 'O010' AND e.numero = 2 "
                    "AND e.id_ouvrier = %s", refuse)
    assert exclu == [("EXCLU",)]
    # Les autres propositions de l'itération 1 ont été annulées (R-RECALC), pas perdues.
    assert q(base, "SELECT count(*) FROM proposal WHERE numero = 1 AND statut = 'ANNULEE' "
                   "AND motif_annulation = 'RECALCUL'")[0][0] == 7
    sain(base)


def test_utilisateur_non_autorise_ou_inconnu(base):
    s.lancer(base.url, J, "L1", "testeur")
    pid, _ = actives(base)["O002"]
    with pytest.raises(s.Interdit) as e:
        s.accepter(base.url, pid, "lecteur")
    assert e.value.code == "ROLE_NON_AUTORISE"
    with pytest.raises(s.Interdit):
        s.accepter(base.url, pid, "inconnu")
    with pytest.raises(s.Interdit):
        s.accepter(base.url, pid, None)
    assert q(base, "SELECT count(*) FROM replacement_assignment") == [(0,)]


def test_acceptation_refusee_si_les_donnees_ont_change(base):
    s.lancer(base.url, J, "L1", "testeur")
    pid, _ = actives(base)["O002"]
    with base.owner() as c:
        c.execute("UPDATE data_version SET version = version + 1 WHERE date = %s", (J,))
    with pytest.raises(s.RefusMetier) as e:
        s.accepter(base.url, pid, "testeur")
    assert e.value.code == "DONNEES_MODIFIEES"


@pytest.mark.parametrize("presents, engages_chaine, code", [(1, 0, "P1"), (9, 9, "P3")])
def test_acceptation_revalide_P1_et_P3(base, monkeypatch, presents, engages_chaine, code):
    s.lancer(base.url, J, "L1", "testeur")
    pid, _ = actives(base)["O002"]
    monkeypatch.setattr(s, "_restants_actuels", lambda *a, **k: (presents, 0, engages_chaine))
    with pytest.raises(s.RefusMetier) as e:
        s.accepter(base.url, pid, "testeur")
    assert e.value.code == code
    assert q(base, "SELECT count(*) FROM replacement_assignment") == [(0,)]


def test_laisser_non_pourvu(base):
    s.lancer(base.url, J, "L1", "testeur")
    besoin = q(base, "SELECT id FROM replacement_request WHERE id_ouvrier_absent = 'O008'")[0][0]
    s.laisser_non_pourvu(base.url, besoin, "testeur", "poste fermé ce jour")
    assert q(base, "SELECT statut_cycle FROM replacement_request WHERE id = %s", besoin) == \
        [("NON_POURVU",)]
    assert "O008" not in actives(base)
    sain(base)


# --- Annulation, clôture ----------------------------------------------------------------------

def test_annuler_un_remplacement_avec_exclusion(base):
    s.lancer(base.url, J, "L1", "testeur")
    pid, ouvrier = actives(base)["O002"]
    id_rempl = s.accepter(base.url, pid, "testeur")
    s.annuler_remplacement(base.url, id_rempl, "testeur", "erreur", exclure=True)
    assert q(base, "SELECT statut FROM replacement_assignment WHERE id = %s", id_rempl) == [("ANNULE",)]
    assert actives(base)["O002"][1] != ouvrier
    with pytest.raises(s.RefusMetier) as e:
        s.annuler_remplacement(base.url, id_rempl, "testeur", "encore", exclure=False)
    assert e.value.code == "REMPLACEMENT_NON_ACTIF"
    sain(base)


def test_cloture_puis_nouveau_calcul_sans_recreer_les_besoins_pourvus(base):
    id_run = s.lancer(base.url, J, "L1", "testeur")
    pid, _ = actives(base)["O002"]
    s.accepter(base.url, pid, "testeur")
    with pytest.raises(s.RefusMetier):
        s.cloturer(base.url, id_run, "testeur", "")
    s.cloturer(base.url, id_run, "testeur", "fin de poste")
    assert actives(base) == {}
    assert q(base, "SELECT count(*) FROM worker_day_slot WHERE statut = 'COURANT' "
                   "AND type = 'RESERVE'") == [(0,)]
    assert q(base, "SELECT count(*) FROM replacement_assignment WHERE statut = 'ACTIF'") == [(1,)]
    sain(base)
    second = s.lancer(base.url, J, "L1", "testeur")
    absents = {a for (a,) in q(base, "SELECT id_ouvrier_absent FROM replacement_request "
                                     "WHERE id_run = %s", second)}
    assert "O002" not in absents and len(absents) == 9   # I11 : O002 déjà pourvu
    sain(base)


# --- Réimport des données ---------------------------------------------------------------------

def _classeur_modifie(chemin, absent_supplementaire):
    """Copie du classeur de référence où un ouvrier présent devient absent (valeurs figées)."""
    wb = load_workbook(CLASSEUR, data_only=True)
    aff = wb["Affectations"]
    entetes = [c.value for c in aff[1]]
    i_id, i_statut = entetes.index("ID_Ouvrier"), entetes.index("Statut_affectation")
    for ligne in aff.iter_rows(min_row=2):
        if ligne[i_id].value == absent_supplementaire:
            ligne[i_statut].value = "Absent"
            nom, chaine, poste = ligne[2].value, ligne[3].value, ligne[5].value
    wb["Absences"].append([dt.datetime(2026, 10, 12), absent_supplementaire, nom, chaine, poste,
                           "Absent", "Absence imprévue", 1])
    wb.save(chemin)
    return chemin


def test_reimport_annule_les_propositions_et_met_a_revoir_le_remplacement(base, tmp_path):
    id_run = s.lancer(base.url, J, "L1", "testeur")
    pid, ouvrier = actives(base)["O002"]
    id_rempl = s.accepter(base.url, pid, "testeur")
    autre, _ = actives(base)["O010"]

    rap = import_journalier.importer(base.url, _classeur_modifie(tmp_path / "m.xlsx", ouvrier),
                                     "testeur", base.archives, True)
    assert rap.statut == "APPLIQUE"
    assert "REMPLACEMENTS_A_REVOIR" in rap.codes()
    assert q(base, "SELECT statut, cause_statut FROM allocation_run WHERE id = %s", id_run) == \
        [("A_RECALCULER", "DONNEES_MODIFIEES")]
    assert q(base, "SELECT count(*) FROM proposal WHERE statut = 'ACTIVE'") == [(0,)]
    assert q(base, "SELECT DISTINCT motif_annulation, id_lot_cause IS NOT NULL FROM proposal "
                   "WHERE statut = 'ANNULEE'") == [("DONNEES_MODIFIEES", True)]
    # Le remplacement n'est PAS annulé automatiquement : il passe à revoir.
    assert q(base, "SELECT statut, statut_revue, cause_revue FROM replacement_assignment") == \
        [("ACTIF", "A_REVOIR", "REMPLACANT_INDISPONIBLE")]
    with pytest.raises(s.RefusMetier) as e:
        s.accepter(base.url, autre, "testeur")
    assert e.value.code == "PROPOSITION_NON_ACTIVE"
    with pytest.raises(s.RefusMetier) as e:
        s.confirmer_apres_reimport(base.url, id_rempl, "testeur", "vérifié")
    assert e.value.code == "NON_CONFORME"
    sain(base)
    s.recalculer(base.url, id_run, "testeur")
    assert q(base, "SELECT statut FROM allocation_run WHERE id = %s", id_run) == [("OUVERT",)]
    sain(base)
    s.annuler_remplacement(base.url, id_rempl, "testeur", "remplaçant absent", exclure=True)
    sain(base)


def test_reimport_sans_effet_maintient_le_remplacement_conforme(base, tmp_path):
    s.lancer(base.url, J, "L1", "testeur")
    pid, _ = actives(base)["O002"]
    s.accepter(base.url, pid, "testeur")
    # Nouvel import d'un autre fichier : un ouvrier d'une autre chaîne devient absent.
    rap = import_journalier.importer(base.url, _classeur_modifie(tmp_path / "m.xlsx", "O149"),
                                     "testeur", base.archives, True)
    assert rap.statut == "APPLIQUE"
    assert q(base, "SELECT statut_revue FROM replacement_assignment") == [("OK",)]
    sain(base)


# --- Concurrence (connexions réelles simultanées) -------------------------------------------

def _en_parallele(*fonctions):
    depart = threading.Barrier(len(fonctions))
    resultats = [None] * len(fonctions)

    def lancer(i, fn):
        depart.wait()
        try:
            resultats[i] = ("ok", fn())
        except s.RefusMetier as exc:
            resultats[i] = ("refus", exc.code)

    fils = [threading.Thread(target=lancer, args=(i, f)) for i, f in enumerate(fonctions)]
    for t in fils:
        t.start()
    for t in fils:
        t.join(60)
    return resultats


@pytest.mark.parametrize("repetition", range(3))
def test_concurrence_deux_acceptations_de_la_meme_proposition(base, repetition):
    s.lancer(base.url, J, "L1", "testeur")
    pid, _ = actives(base)["O002"]
    res = _en_parallele(lambda: s.accepter(base.url, pid, "testeur"),
                        lambda: s.accepter(base.url, pid, "testeur"))
    assert sorted(r[0] for r in res) == ["ok", "refus"]
    assert q(base, "SELECT count(*) FROM replacement_assignment") == [(1,)]
    sain(base)


@pytest.mark.parametrize("repetition", range(3))
def test_concurrence_acceptation_contre_refus(base, repetition):
    s.lancer(base.url, J, "L1", "testeur")
    pid, _ = actives(base)["O002"]
    res = _en_parallele(lambda: s.accepter(base.url, pid, "testeur"),
                        lambda: s.refuser(base.url, pid, "testeur", "refus concurrent"))
    assert sorted(r[0] for r in res) == ["ok", "refus"]
    statut = q(base, "SELECT statut FROM proposal WHERE id = %s", pid)[0][0]
    assert statut in ("ACCEPTEE", "REFUSEE")
    assert q(base, "SELECT count(*) FROM replacement_assignment")[0][0] == (statut == "ACCEPTEE")
    sain(base)


@pytest.mark.parametrize("repetition", range(3))
def test_concurrence_deux_calculs_sur_deux_chaines_sans_double_reservation(base, repetition):
    res = _en_parallele(lambda: s.lancer(base.url, J, "L1", "testeur"),
                        lambda: s.lancer(base.url, J, "L2", "testeur"))
    assert [r[0] for r in res] == ["ok", "ok"]
    ouvriers = [o for (o,) in q(base, "SELECT id_ouvrier FROM proposal WHERE statut = 'ACTIVE'")]
    assert len(ouvriers) == len(set(ouvriers))
    # R-CASCADE (profil provisoire) : aucun ouvrier de L1 proposé à L2, ni l'inverse.
    chaines = q(base, "SELECT DISTINCT r.id_chaine_cible, e.chaine_donneuse FROM proposal p "
                      "JOIN candidate_evaluation e ON e.id = p.id_evaluation "
                      "JOIN replacement_request r ON r.id = p.id_besoin WHERE p.statut = 'ACTIVE'")
    assert all(cible != donneuse and donneuse not in ("L1", "L2") for cible, donneuse in chaines)
    sain(base)


def test_concurrence_deux_lancements_identiques(base):
    res = _en_parallele(lambda: s.lancer(base.url, J, "L1", "testeur"),
                        lambda: s.lancer(base.url, J, "L1", "testeur"))
    assert sorted(r[0] for r in res) == ["ok", "refus"]
    assert q(base, "SELECT count(*) FROM allocation_run") == [(1,)]
    sain(base)


def test_concurrence_acceptations_et_import(base, tmp_path):
    s.lancer(base.url, J, "L1", "testeur")
    props = list(actives(base).values())
    fichier = _classeur_modifie(tmp_path / "m.xlsx", "O149")
    res = _en_parallele(
        lambda: s.accepter(base.url, props[0][0], "testeur"),
        lambda: s.accepter(base.url, props[1][0], "testeur"),
        lambda: import_journalier.importer(base.url, fichier, "testeur", base.archives, True).statut)
    assert res[2] == ("ok", "APPLIQUE")
    sain(base)
    # Chaque acceptation a soit réussi avant l'import, soit été refusée (proposition annulée).
    for r in res[:2]:
        assert r[0] == "ok" or r[1] in ("PROPOSITION_NON_ACTIVE", "RUN_NON_OUVERT",
                                        "DONNEES_MODIFIEES")


# --- Garanties de la base et de la configuration --------------------------------------------

def test_lignes_d_historique_non_modifiables(base):
    s.lancer(base.url, J, "L1", "testeur")
    pid, _ = actives(base)["O002"]
    s.accepter(base.url, pid, "testeur")
    with base.owner() as c:
        for sql in ("UPDATE allocation_decision SET motif = 'x'", "DELETE FROM allocation_decision",
                    "UPDATE candidate_evaluation SET rang = 1", "DELETE FROM allocation_iteration",
                    "UPDATE proposal SET statut = 'ACTIVE' WHERE statut = 'ACCEPTEE'"):
            with pytest.raises(psycopg.errors.RestrictViolation):
                c.execute(sql)
            c.rollback()


def test_configuration_active_immuable_et_non_validee(base):
    with base.owner() as c:
        with pytest.raises(psycopg.errors.RestrictViolation):
            c.execute("UPDATE allocation_configuration SET p3 = 99 WHERE statut = 'ACTIVE'")
        c.rollback()
        indicateurs = c.execute(
            "SELECT m_bloc_valide, m_alerte_valide, p3_valide, priorites_valide FROM "
            "allocation_configuration WHERE statut = 'ACTIVE'").fetchone()
    assert indicateurs == (False, False, False, False)


def test_hors_TEST_aucune_configuration_provisoire_ne_peut_etre_activee(base_migree):
    with base_migree.owner() as c:
        c.execute("INSERT INTO environment (usage, libelle) VALUES ('EXPLOITATION', 'x')")
        c.execute("INSERT INTO app_user (identifiant, nom, role) VALUES ('u', 'U', 'R')")
        c.execute("INSERT INTO import_batch (type, perimetre, fichier_source, empreinte_fichier, "
                  "auteur, statut, date_application) VALUES ('IMPORT', 'REFERENTIEL', 'x', "
                  "repeat('a', 64), 1, 'APPLIQUE', now())")
        c.execute("INSERT INTO production_line VALUES ('L1', 'P', 1, 1, 'Moyenne', 1, 1)")
        v = c.execute("INSERT INTO allocation_configuration (libelle, m_bloc, m_alerte, p3, "
                      "mode_expiration, auteur) VALUES ('x', 1, 1, 1, 'MANUELLE', 1) "
                      "RETURNING version").fetchone()[0]
        c.execute("INSERT INTO config_priority_rank VALUES (%s, 'Moyenne', 1)", (v,))
        with pytest.raises(psycopg.errors.CheckViolation, match="hors TEST"):
            c.execute("UPDATE allocation_configuration SET statut = 'ACTIVE', active_par = 1, "
                      "date_activation = now() WHERE version = %s", (v,))
    with pytest.raises(Exception, match="TEST"):
        configuration.installer_profil_provisoire(base_migree.owner_url, "A", "u")
